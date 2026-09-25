# -*- coding: utf-8 -*-
"""Export User_Manual.md to Word (.docx) and PDF."""
from pathlib import Path
import re

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, Inches
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    Preformatted,
)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

BASE = Path(__file__).resolve().parent
MD_PATH = BASE / "User_Manual.md"
DOCX_PATH = BASE / "User_Manual.docx"
PDF_PATH = BASE / "User_Manual.pdf"
FONT_PATH = Path(r"C:\Windows\Fonts\simhei.ttf")


def set_run_font(run, name="Microsoft YaHei", size=None, bold=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_runs(paragraph, content, bold=False, size=None):
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", content)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name="Consolas", size=size or 10, bold=bold)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, bold=bold)


def build_docx(text: str) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    lines = text.splitlines()
    i = 0
    in_code = False
    code_buf = []
    table_buf = []

    def flush_table():
        nonlocal table_buf
        if not table_buf:
            return
        rows = []
        for row in table_buf:
            if re.match(r"^\|?\s*-+", row):
                continue
            cells = [c.strip() for c in row.strip().strip("|").split("|")]
            rows.append(cells)
        if rows:
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r_idx, row in enumerate(rows):
                for c_idx, cell in enumerate(row):
                    clean = re.sub(r"\*\*|`", "", cell)
                    para = table.rows[r_idx].cells[c_idx].paragraphs[0]
                    para.text = ""
                    run = para.add_run(clean)
                    set_run_font(run, size=10, bold=(r_idx == 0))
        table_buf = []

    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_buf))
                set_run_font(run, name="Consolas", size=10)
                code_buf = []
                in_code = False
            else:
                flush_table()
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue
        if line.startswith("|"):
            table_buf.append(line)
            i += 1
            continue
        flush_table()

        if line.strip() == "---":
            i += 1
            continue
        if line.startswith("# "):
            p = doc.add_heading("", level=0)
            add_runs(p, line[2:].strip(), size=18)
        elif line.startswith("## "):
            p = doc.add_heading("", level=1)
            add_runs(p, line[3:].strip(), size=14)
        elif line.startswith("### "):
            p = doc.add_heading("", level=2)
            add_runs(p, line[4:].strip(), size=12)
        elif line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            add_runs(p, line[2:].strip())
            for run in p.runs:
                run.italic = True
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\d+\. ", "", line))
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, line[2:])
        elif line.strip() == "":
            pass
        else:
            p = doc.add_paragraph()
            add_runs(p, line)
        i += 1
    flush_table()
    doc.save(DOCX_PATH)


def md_inline(s: str) -> str:
    s = s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    s = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", s)
    s = re.sub(r"`([^`]+)`", r"<font face='Courier' size='9'>\1</font>", s)
    return s


def build_pdf(text: str) -> None:
    if not FONT_PATH.exists():
        raise FileNotFoundError(f"Chinese font not found: {FONT_PATH}")
    pdfmetrics.registerFont(TTFont("SimHei", str(FONT_PATH)))

    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="CNTitle", fontName="SimHei", fontSize=16, leading=22, spaceAfter=10
        )
    )
    styles.add(
        ParagraphStyle(
            name="CNH1",
            fontName="SimHei",
            fontSize=13,
            leading=18,
            spaceBefore=12,
            spaceAfter=6,
        )
    )
    styles.add(
        ParagraphStyle(
            name="CNH2",
            fontName="SimHei",
            fontSize=11,
            leading=16,
            spaceBefore=8,
            spaceAfter=4,
        )
    )
    styles.add(
        ParagraphStyle(
            name="CNBody", fontName="SimHei", fontSize=10, leading=15, spaceAfter=4
        )
    )
    styles.add(
        ParagraphStyle(
            name="CNQuote",
            fontName="SimHei",
            fontSize=9.5,
            leading=14,
            leftIndent=10,
            textColor=colors.HexColor("#333333"),
            spaceAfter=6,
        )
    )
    styles.add(
        ParagraphStyle(
            name="CNCode",
            fontName="Courier",
            fontSize=9,
            leading=12,
            backColor=colors.HexColor("#f4f4f4"),
            spaceBefore=4,
            spaceAfter=6,
            leftIndent=6,
            rightIndent=6,
        )
    )
    styles.add(ParagraphStyle(name="CNCell", fontName="SimHei", fontSize=9, leading=12))

    story = []
    lines = text.splitlines()
    i = 0
    in_code = False
    code_buf = []
    table_buf = []

    def flush_pdf_table():
        nonlocal table_buf
        if not table_buf:
            return
        rows = []
        for row in table_buf:
            if re.match(r"^\|?\s*-+", row):
                continue
            cells = [c.strip() for c in row.strip().strip("|").split("|")]
            rows.append([Paragraph(md_inline(c), styles["CNCell"]) for c in cells])
        if rows:
            ncols = len(rows[0])
            if ncols == 3:
                widths = [50 * mm, 35 * mm, 85 * mm]
            elif ncols == 2:
                widths = [55 * mm, 115 * mm]
            else:
                widths = [170 * mm / ncols] * ncols
            t = Table(rows, hAlign="LEFT", colWidths=widths)
            t.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
                        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 4),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                        ("TOPPADDING", (0, 0), (-1, -1), 3),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                    ]
                )
            )
            story.append(t)
            story.append(Spacer(1, 6))
        table_buf = []

    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            if in_code:
                story.append(Preformatted("\n".join(code_buf), styles["CNCode"]))
                code_buf = []
                in_code = False
            else:
                flush_pdf_table()
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue
        if line.startswith("|"):
            table_buf.append(line)
            i += 1
            continue
        flush_pdf_table()

        if line.strip() == "---":
            i += 1
            continue
        if line.startswith("# "):
            story.append(Paragraph(md_inline(line[2:].strip()), styles["CNTitle"]))
        elif line.startswith("## "):
            story.append(Paragraph(md_inline(line[3:].strip()), styles["CNH1"]))
        elif line.startswith("### "):
            story.append(Paragraph(md_inline(line[4:].strip()), styles["CNH2"]))
        elif line.startswith("> "):
            story.append(Paragraph(md_inline(line[2:].strip()), styles["CNQuote"]))
        elif re.match(r"^\d+\. ", line):
            story.append(Paragraph(md_inline(line), styles["CNBody"]))
        elif line.startswith("- "):
            story.append(Paragraph("• " + md_inline(line[2:]), styles["CNBody"]))
        elif line.strip() == "":
            story.append(Spacer(1, 4))
        else:
            story.append(Paragraph(md_inline(line), styles["CNBody"]))
        i += 1
    flush_pdf_table()

    pdf = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title="Temperature Dashboard User Manual",
    )
    pdf.build(story)


def main():
    if not MD_PATH.exists():
        raise FileNotFoundError(MD_PATH)
    text = MD_PATH.read_text(encoding="utf-8")
    build_docx(text)
    print("Wrote", DOCX_PATH)
    build_pdf(text)
    print("Wrote", PDF_PATH)
    print("DOCX bytes:", DOCX_PATH.stat().st_size)
    print("PDF bytes:", PDF_PATH.stat().st_size)


if __name__ == "__main__":
    main()
